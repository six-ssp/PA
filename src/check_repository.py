"""只读检查已发布成果；--with-intermediate 额外逐单元格核对本地求解结果。"""

from __future__ import annotations

import argparse
import ast
from bisect import bisect_right
import csv
import json
import math
import re
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path
from urllib.parse import unquote, urlsplit

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def radius_from_attachment(time_s: float, samples: list[tuple[float, float]]) -> float:
    """独立按附件的 cm 单位插值，避免用被检查的模型实现验证自身。"""
    index = bisect_right([row[0] for row in samples], time_s)
    if index == 0:
        return samples[0][1]
    if index == len(samples):
        return samples[-1][1]
    t0, r0 = samples[index - 1]
    t1, r1 = samples[index]
    return r0 + (r1 - r0) * (time_s - t0) / (t1 - t0)


def require(condition: bool, message: str) -> None:
    """不用 assert，确保 python -O 也不会跳过成果检查。"""
    if not condition:
        raise ValueError(message)


def read_json(path: Path) -> dict:
    def reject(value: str) -> None:
        raise ValueError(f"{path.name}: 非标准 JSON 数值 {value}")
    return json.loads(path.read_text(encoding="utf-8"), parse_constant=reject)


def close(a: float, b: float, label: str, tolerance: float = 1e-7) -> None:
    require(math.isfinite(a) and math.isfinite(b) and abs(a - b) <= tolerance,
            f"{label}: {a} != {b}")


def check_tables(summary: dict, with_intermediate: bool) -> int:
    """逐行核对时间、有限性、收缩后的空白区域；可选与 JSON 全量比较。"""
    count = 0
    radius_book = load_workbook(ROOT / "附件2.xlsx", read_only=True, data_only=True)
    try:
        radius_samples = [(float(row[0]), float(row[1])) for row in radius_book.worksheets[0].iter_rows(min_row=2, values_only=True) if row[0] is not None]
    finally:
        radius_book.close()
    for number in range(1, 5):
        name = f"result{number}"
        payload = read_json(ROOT / "intermediate" / f"{name}.json") if with_intermediate else None
        end = {1: 1800, 2: 10800}.get(number, summary.get(f"problem{number}_excel_end_s"))
        step = 1 if number < 3 else 60
        specs = [("温度", "temperature"), ("水分浓度", "moisture")] if number < 3 else [("Sheet1", "moisture")]
        book = load_workbook(ROOT / "results" / f"{name}.xlsx", read_only=True, data_only=False)
        try:
            require(book.sheetnames == [s[0] for s in specs], f"{name}: 工作表名称不符")
            if payload:
                require(payload["time"] == list(range(0, end + 1, step)), f"{name}: JSON 时间网格不符")
                if number >= 3:
                    close(payload["drying_time_s"], summary[f"problem{number}_drying_time_s"], f"{name}: 事件时间")
            for sheet_name, key in specs:
                sheet = book[sheet_name]
                rows = sheet.iter_rows(values_only=True)
                header = list(next(rows))
                distances = [round(i / 10, 1) for i in range(21)]
                if number == 4:
                    distances.append("药材表面")
                require(header[1:] == distances, f"{name}/{sheet_name}: 位置表头不符")
                expected_rows = end // step + 1
                if payload:
                    require(payload["distance"] == distances, f"{name}: JSON 位置不符")
                    require(len(payload[key]) == expected_rows, f"{name}: JSON 行数不符")
                # 半径来自原始附件，不用输出表自身推断其空白区是否正确。
                row_count = 0
                for i, row in enumerate(rows):
                    require(i < expected_rows, f"{name}: 多余数据行")
                    close(row[0], i * step, f"{name}: 第 {i + 2} 行时间", 0)
                    if payload:
                        require(len(payload[key][i]) == len(distances), f"{name}: JSON 列数不符")
                    radius_cm = radius_from_attachment(i * step, radius_samples) if number == 4 else 2.0
                    for j, value in enumerate(row[1:]):
                        outside = number == 4 and j < 21 and distances[j] > radius_cm + 1e-10
                        if outside:
                            require(value is None, f"{name}: 药材外部应留空，行 {i + 2} 列 {j + 2}")
                        else:
                            require(isinstance(value, (float, int)) and math.isfinite(value), f"{name}: 非法场值")
                            require(value >= 0, f"{name}: 负场值")
                            close(value, round(value, 4), f"{name}: 场值未按四位小数输出", 1e-9)
                            if i == 0:
                                close(value, 28 if key == "temperature" else 2.55, f"{name}: 初值")
                        if payload:
                            expected = payload[key][i][j]
                            if value is None or expected is None:
                                require(value is expected, f"{name}: JSON/Excel 空值不一致")
                            else:
                                close(value, expected, f"{name}: JSON/Excel 行 {i + 2} 列 {j + 2}", 1e-10)
                        count += 1
                    row_count += 1
                require(row_count == expected_rows, f"{name}: Excel 行数不符")
        finally:
            book.close()
    return count


def check_metrics(summary: dict) -> None:
    grid = summary["grid_convergence"]
    with (ROOT / "results/grid_convergence.csv").open(encoding="utf-8-sig", newline="") as handle:
        csv_grid = list(csv.DictReader(handle))
    require([int(row["nodes"]) for row in csv_grid] == [81, 161, 321, 641], "收敛网格档位不符")
    require(len(grid) == len(csv_grid), "收敛摘要长度不符")
    for a, b in zip(grid, csv_grid):
        require(a["nodes"] == int(b["nodes"]), "收敛节点数不一致")
        for q in (3, 4):
            close(a[f"problem{q}_time_h"], float(b[f"problem{q}_time_h"]), "收敛 CSV/摘要")
    errors = read_json(ROOT / "results/error_analysis.json")
    require(errors["spatial"]["nodes"] == [row["nodes"] for row in grid], "误差网格不一致")
    sensitivity = read_json(ROOT / "results/sensitivity_summary.json")
    for q in (3, 4):
        key = f"problem{q}"
        t = summary[f"{key}_drying_time_s"]
        close(t / 3600, summary[f"{key}_drying_time_h"], f"{key}: 小时换算")
        close(t / 3600, grid[-1][f"{key}_time_h"], f"{key}: 正式网格")
        require(summary[f"{key}_excel_end_s"] == math.ceil(t / 60) * 60, f"{key}: Excel 终点")
        require(len(errors["spatial"][f"{key}_time_h"]) == len(grid), "误差数据长度不符")
        for a, b in zip(errors["spatial"][f"{key}_time_h"], grid):
            close(a, b[f"{key}_time_h"], f"{key}: 误差数据")
        close(sensitivity["baseline_drying_time_h"][key], grid[2][f"{key}_time_h"], f"{key}: 敏感度基准")
    fixed = summary["appendix4_fixed_radius_drying_time_s"]
    shrink = summary["problem4_drying_time_s"]
    close(100 * (fixed - shrink) / fixed, summary["appendix4_shrinkage_reduction_percent"], "独立收缩对照")
    with (ROOT / "results/sensitivity_sweep.csv").open(encoding="utf-8-sig", newline="") as handle:
        sweep = list(csv.DictReader(handle))
    require(len(sweep) == sensitivity["sweep_row_count"] == 196, "敏感度扫描记录不完整")
    for q in ("problem3", "problem4"):
        for parameter in sensitivity["local_sensitivity"][q]:
            group = [r for r in sweep if r["problem"] == q and r["parameter"] == parameter]
            points = sensitivity["temperature_offset_points" if parameter == "stable_temperature" else "scale_factor_points"]
            require(sorted(float(r["value"]) for r in group) == points, f"{q}/{parameter}: 缺失或重复取点")
            for row in group:
                t = float(row["drying_time_h"])
                require(t > 0 and math.isfinite(t), "敏感度时长非法")
                baseline = sensitivity["baseline_drying_time_h"][q]
                close(100 * (t / baseline - 1), float(row["change_from_baseline_percent"]), "敏感度百分比")
    event = read_json(ROOT / "figures/method_figures_validation.json")
    close(event["drying_time_s"], summary["problem3_drying_time_s"], "图 15 事件时间", 1.0)
    close(event["event_residual"], 0, "图 15 事件残差", 2e-5)


def check_sources_and_figures() -> None:
    # 只检查版本库所属文件，不扫描用户的其他实验目录或私人论文。
    raw = subprocess.check_output(["git", "ls-files", "-z"], cwd=ROOT)
    tracked = [Path(p) for p in raw.decode("utf-8").split("\0") if p]
    for relative in tracked:
        path = ROOT / relative
        if path.suffix in {".py", ".md", ".json", ".csv", ".mjs"}:
            text = path.read_text(encoding="utf-8-sig")
            require("\ufffd" not in text, f"{relative}: 含编码替代字符")
            if path.suffix == ".py":
                ast.parse(text, filename=str(relative))
            if path.suffix == ".json":
                read_json(path)
            if path.suffix == ".md":
                # 简单检查本地 Markdown 链接；不联网，也不把远程链接判作丢失文件。
                for target in re.findall(r"!?\[[^\]]*\]\(([^\s)]+)\)", text):
                    url = urlsplit(target.strip("<>"))
                    if url.scheme or not url.path:
                        continue
                    require((path.parent / unquote(url.path)).exists(), f"{relative}: 失效链接 {target}")
                math_blocks = re.findall(r"^\$\$\s*\n(.*?)^\$\$", text, flags=re.M | re.S)
                for block in math_blocks:
                    require(not re.search(r"(?<!\\)%", block), f"{relative}: 数学环境中百分号必须转义")
    for number in range(1, 16):
        pngs = list((ROOT / "figures").glob(f"{number:02d}_*.png"))
        require(len(pngs) == 1, f"图 {number}: PNG 缺失或重名")
        require(pngs[0].read_bytes()[:8] == b"\x89PNG\r\n\x1a\n", f"图 {number}: PNG 文件损坏")
        require(ET.parse(pngs[0].with_suffix(".svg")).getroot().tag.endswith("}svg"), f"图 {number}: SVG 非法")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--with-intermediate", action="store_true", help="要求四份本地 JSON 并逐单元格比较 Excel")
    parser.add_argument("--tables-only", action="store_true", help="仅检查四问表格，供重新生成分析数据之前使用")
    args = parser.parse_args()
    summary = read_json(ROOT / "intermediate/summary.json")
    if not args.tables_only:
        check_sources_and_figures()
        check_metrics(summary)
    count = check_tables(summary, args.with_intermediate)
    scope = "四问表格" if args.tables_only else "源码/文档、15 组 PNG/SVG、摘要/收敛/196 条敏感度记录"
    print(f"PASS: {scope}、{count:,} 个场值检查通过")
    print("范围：成果一致性与格式检查；不替代独立实验验证或全量重新求解。")


if __name__ == "__main__":
    main()
