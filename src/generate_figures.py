"""从正式结果生成论文级多面板图，并输出合理性诊断指标。"""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import load_workbook

from model import load_environment_xlsx, load_radius_xlsx


ROOT = Path(__file__).resolve().parents[1]
INTERMEDIATE = ROOT / "intermediate"
RESULTS = ROOT / "results"
FIGURES = ROOT / "figures"

COLORS = {
    "navy": "#17324D",
    "blue": "#2878B5",
    "cyan": "#36B5C5",
    "teal": "#1B998B",
    "gold": "#E6A23C",
    "orange": "#E76F51",
    "red": "#C44536",
    "purple": "#7E57C2",
    "ink": "#25313C",
    "muted": "#697786",
    "grid": "#DCE3EA",
    "paper": "#F7F9FC",
}


def configure_style() -> None:
    """统一论文图的字体、线宽、网格和导出质量。"""
    mpl.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "figure.facecolor": COLORS["paper"],
        "axes.facecolor": "white",
        "axes.edgecolor": COLORS["grid"],
        "axes.labelcolor": COLORS["ink"],
        "axes.titlecolor": COLORS["ink"],
        "xtick.color": COLORS["muted"],
        "ytick.color": COLORS["muted"],
        "text.color": COLORS["ink"],
        "grid.color": COLORS["grid"],
        "grid.linewidth": 0.7,
        "grid.alpha": 0.72,
        "axes.titleweight": "bold",
        "axes.titlesize": 12,
        "axes.labelsize": 10,
        "legend.frameon": False,
        "legend.fontsize": 8.5,
        "lines.linewidth": 2.0,
        "savefig.facecolor": COLORS["paper"],
        "savefig.bbox": "tight",
    })


def read_xlsx_sheet(path: Path, sheet_name: str) -> tuple[list, np.ndarray, np.ndarray]:
    """只读提取正式 Excel 中的一张结果表。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))[1:]
    time_values: list[float] = []
    field_values: list[list[float]] = []
    for row in rows:
        if row[0] is None:
            continue
        time_values.append(float(row[0]))
        field_values.append([np.nan if value is None else float(value) for value in row[1:1 + len(header)]])
    workbook.close()
    return header, np.asarray(time_values), np.asarray(field_values)


def load_payloads() -> dict[str, dict]:
    """优先使用 run_all.py 的中间 JSON；缺失时回退到已跟踪的正式 Excel。"""
    names = ["result1", "result2", "result3", "result4"]
    if all((INTERMEDIATE / f"{name}.json").exists() for name in names):
        return {
            name: json.loads((INTERMEDIATE / f"{name}.json").read_text(encoding="utf-8"))
            for name in names
        }

    result: dict[str, dict] = {}
    for name in ("result1", "result2"):
        header_t, times_t, temperature = read_xlsx_sheet(RESULTS / f"{name}.xlsx", "温度")
        header_c, times_c, moisture = read_xlsx_sheet(RESULTS / f"{name}.xlsx", "水分浓度")
        if header_t != header_c or not np.array_equal(times_t, times_c):
            raise RuntimeError(f"{name}.xlsx 的温度、水分表结构不一致")
        result[name] = {
            "time": times_t.tolist(),
            "distance": header_t,
            "temperature": temperature.tolist(),
            "moisture": moisture.tolist(),
        }
    for name in ("result3", "result4"):
        header, times, moisture = read_xlsx_sheet(RESULTS / f"{name}.xlsx", "Sheet1")
        result[name] = {
            "time": times.tolist(),
            "distance": header,
            "moisture": moisture.tolist(),
        }
    return result


def clean_axis(axis: plt.Axes, *, grid: bool = True) -> None:
    axis.spines[["top", "right"]].set_visible(False)
    axis.spines[["left", "bottom"]].set_color(COLORS["grid"])
    if grid:
        axis.grid(True, linestyle="-", zorder=0)
    axis.set_axisbelow(True)


def panel_label(axis: plt.Axes, label: str) -> None:
    axis.text(
        -0.10, 1.06, label, transform=axis.transAxes,
        fontsize=12, fontweight="bold", color=COLORS["navy"], va="bottom",
    )


def nearest_indices(times: np.ndarray, requested: list[float]) -> list[int]:
    return [int(np.argmin(np.abs(times - value))) for value in requested]


def save_figure(figure: plt.Figure, stem: str) -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    figure.savefig(FIGURES / f"{stem}.png", dpi=220)
    svg_path = FIGURES / f"{stem}.svg"
    figure.savefig(svg_path)
    # Windows 下 Matplotlib 默认写 CRLF；统一为 LF，避免 Git 把 CR 识别成行尾空白。
    svg_content = "\n".join(
        line.rstrip() for line in svg_path.read_text(encoding="utf-8").splitlines()
    ) + "\n"
    with svg_path.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(svg_content)
    plt.close(figure)


def plot_short_term_profiles(data: dict[str, dict]) -> None:
    figure, axes = plt.subplots(2, 2, figsize=(13.2, 8.6), constrained_layout=True)
    specs = [
        ("result1", "temperature", [0, 300, 600, 1200, 1800], "问题 1 · 温度径向演化", "温度 / °C"),
        ("result1", "moisture", [0, 300, 600, 1200, 1800], "问题 1 · 含水率径向演化", "含水率 / (kg/kg)"),
        ("result2", "temperature", [0, 1800, 3600, 7200, 10800], "问题 2 · 温度径向演化", "温度 / °C"),
        ("result2", "moisture", [0, 1800, 3600, 7200, 10800], "问题 2 · 含水率径向演化", "含水率 / (kg/kg)"),
    ]
    palette = [COLORS["navy"], COLORS["blue"], COLORS["cyan"], COLORS["gold"], COLORS["red"]]
    for letter, axis, (name, field, requested, title, ylabel) in zip("ABCD", axes.flat, specs):
        times = np.asarray(data[name]["time"], dtype=float)
        radius = np.asarray(data[name]["distance"], dtype=float)
        values = np.asarray(data[name][field], dtype=float)
        for color, index in zip(palette, nearest_indices(times, requested)):
            label = f"{times[index] / 60:.0f} min" if times[index] < 3600 else f"{times[index] / 3600:.1f} h"
            axis.plot(radius, values[index], color=color, marker="o", markersize=3.2, label=label)
        axis.set(title=title, xlabel="距中心径向距离 / cm", ylabel=ylabel)
        axis.legend(ncol=2, loc="best")
        clean_axis(axis)
        panel_label(axis, letter)
    figure.suptitle("短时传热—传质耦合的径向剖面对比", fontsize=16, fontweight="bold", color=COLORS["navy"])
    save_figure(figure, "01_short_term_radial_profiles")


def plot_long_term_comparison(data: dict[str, dict], summary: dict) -> None:
    figure, axes = plt.subplots(2, 2, figsize=(13.2, 8.8), constrained_layout=True)
    q3 = data["result3"]
    q4 = data["result4"]
    time3 = np.asarray(q3["time"], dtype=float) / 3600.0
    time4 = np.asarray(q4["time"], dtype=float) / 3600.0
    c3 = np.asarray(q3["moisture"], dtype=float)
    c4 = np.asarray(q4["moisture"], dtype=float)

    axes[0, 0].plot(time3, c3[:, 0], color=COLORS["blue"], label="问题 3 · 中心")
    axes[0, 0].plot(time3, c3[:, -1], color=COLORS["blue"], linestyle="--", label="问题 3 · 表面")
    axes[0, 0].plot(time4, c4[:, 0], color=COLORS["orange"], label="问题 4 · 中心")
    axes[0, 0].plot(time4, c4[:, -1], color=COLORS["orange"], linestyle="--", label="问题 4 · 动态表面")
    axes[0, 0].axhline(0.15, color=COLORS["red"], linestyle=":", linewidth=1.5, label="烘干阈值")
    axes[0, 0].set(title="中心—表面长时间干燥轨迹", xlabel="时间 / h", ylabel="含水率 / (kg/kg)")
    axes[0, 0].legend(ncol=2)
    clean_axis(axes[0, 0])

    fractions = [0.0, 0.25, 0.50, 0.75, 1.0]
    colors = [COLORS["navy"], COLORS["blue"], COLORS["teal"], COLORS["gold"], COLORS["red"]]
    radius3 = np.asarray(q3["distance"], dtype=float)
    for color, fraction in zip(colors, fractions):
        index = int(round(fraction * (len(time3) - 1)))
        axes[0, 1].plot(radius3, c3[index], color=color, label=f"{fraction:.0%} $t_3$")
    axes[0, 1].set(title="问题 3 · 不同干燥进程的径向剖面", xlabel="距中心距离 / cm", ylabel="含水率 / (kg/kg)")
    axes[0, 1].legend(ncol=2)
    clean_axis(axes[0, 1])

    radius4 = np.asarray(q4["distance"][:-1], dtype=float)
    radius_time, radius_m = load_radius_xlsx(ROOT / "附件2.xlsx")
    for color, fraction in zip(colors, fractions):
        index = int(round(fraction * (len(time4) - 1)))
        current_radius = float(np.interp(time4[index] * 3600.0, radius_time, radius_m) * 100.0)
        mask = np.isfinite(c4[index, :-1])
        x_values = np.append(radius4[mask], current_radius)
        y_values = np.append(c4[index, :-1][mask], c4[index, -1])
        order = np.argsort(x_values)
        axes[1, 0].plot(x_values[order], y_values[order], color=color, marker="o", markersize=2.8, label=f"{fraction:.0%} $t_4$")
    axes[1, 0].set(title="问题 4 · 收缩域内的径向剖面", xlabel="实时物理半径 / cm", ylabel="含水率 / (kg/kg)")
    axes[1, 0].legend(ncol=2)
    clean_axis(axes[1, 0])

    labels = ["问题 3\n附录 3 固定半径", "问题 4\n附录 4 收缩", "附录 4\n固定半径对照"]
    values = [
        summary["problem3_drying_time_h"],
        summary["problem4_drying_time_h"],
        summary["appendix4_fixed_radius_drying_time_h"],
    ]
    bars = axes[1, 1].bar(labels, values, color=[COLORS["blue"], COLORS["orange"], COLORS["purple"]], width=0.64)
    for bar, value in zip(bars, values):
        axes[1, 1].text(bar.get_x() + bar.get_width() / 2, value + 2.2, f"{value:.2f} h", ha="center", va="bottom", fontweight="bold")
    axes[1, 1].annotate(
        f"同物性纯收缩降幅\n{summary['appendix4_shrinkage_reduction_percent']:.2f}%",
        xy=(1, values[1]), xytext=(1.52, 82),
        arrowprops={"arrowstyle": "->", "color": COLORS["orange"]},
        ha="center", color=COLORS["orange"], fontweight="bold",
    )
    axes[1, 1].set(title="结束时间：问题差异与收缩效应分离", ylabel="烘干时间 / h")
    clean_axis(axes[1, 1], grid=False)
    axes[1, 1].grid(axis="y")

    for letter, axis in zip("ABCD", axes.flat):
        panel_label(axis, letter)
    figure.suptitle("长时干燥过程的多尺度对比", fontsize=16, fontweight="bold", color=COLORS["navy"])
    save_figure(figure, "02_long_term_drying_comparison")


def plot_spatiotemporal_fields(data: dict[str, dict]) -> None:
    figure, axes = plt.subplots(1, 2, figsize=(13.2, 5.0), constrained_layout=True, sharey=True)
    cmap = LinearSegmentedColormap.from_list(
        "advanced_moisture",
        ["#F5F7FA", "#B8E0D2", "#4CB5AE", "#2878B5", "#5146A5", "#2A174F"],
    )
    q3 = data["result3"]
    time3 = np.asarray(q3["time"], dtype=float) / 3600.0
    radius3 = np.asarray(q3["distance"], dtype=float)
    c3 = np.asarray(q3["moisture"], dtype=float)
    mesh3 = axes[0].pcolormesh(radius3, time3, c3, cmap=cmap, shading="auto", vmin=0.05, vmax=2.55, rasterized=True)
    axes[0].contour(radius3, time3, c3, levels=[0.15, 0.5, 1.0, 2.0], colors="white", linewidths=0.75, alpha=0.85)
    axes[0].set(title="问题 3 · 固定半径水分场", xlabel="距中心距离 / cm", ylabel="时间 / h")

    q4 = data["result4"]
    time4 = np.asarray(q4["time"], dtype=float) / 3600.0
    radius4 = np.asarray(q4["distance"][:-1], dtype=float)
    c4 = np.asarray(q4["moisture"], dtype=float)[:, :-1]
    mesh4 = axes[1].pcolormesh(radius4, time4, c4, cmap=cmap, shading="auto", vmin=0.05, vmax=2.55, rasterized=True)
    radius_time, radius_m = load_radius_xlsx(ROOT / "附件2.xlsx")
    boundary = np.interp(time4 * 3600.0, radius_time, radius_m) * 100.0
    axes[1].plot(boundary, time4, color="white", linewidth=2.0, label="实时表面 R(t)")
    axes[1].plot(boundary, time4, color=COLORS["red"], linewidth=0.9)
    axes[1].set(title="问题 4 · 收缩物理域水分场", xlabel="固定物理坐标 / cm")
    axes[1].legend(loc="lower right")

    for letter, axis in zip("AB", axes):
        clean_axis(axis, grid=False)
        panel_label(axis, letter)
    colorbar = figure.colorbar(mesh4, ax=axes, orientation="vertical", shrink=0.88, pad=0.025)
    colorbar.set_label("含水率 / (kg/kg)")
    figure.suptitle("含水率的时空演化与收缩边界", fontsize=16, fontweight="bold", color=COLORS["navy"])
    save_figure(figure, "03_spatiotemporal_moisture_fields")


def plot_grid_error(error: dict) -> None:
    spatial = error["spatial"]
    nodes = np.asarray(spatial["nodes"], dtype=float)
    q3 = np.asarray(spatial["problem3_time_h"], dtype=float)
    q4 = np.asarray(spatial["problem4_time_h"], dtype=float)
    e3 = np.asarray(spatial["problem3_abs_error_vs_641_s"], dtype=float)
    e4 = np.asarray(spatial["problem4_abs_error_vs_641_s"], dtype=float)
    figure, axes = plt.subplots(1, 2, figsize=(12.8, 4.8), constrained_layout=True)

    axes[0].plot(nodes, q3, color=COLORS["blue"], marker="o", markersize=6, label="问题 3")
    axes[0].plot(nodes, q4, color=COLORS["orange"], marker="s", markersize=6, label="问题 4")
    for x_value, y_value in zip(nodes, q3):
        axes[0].annotate(f"{y_value:.3f}", (x_value, y_value), xytext=(0, 8), textcoords="offset points", ha="center", fontsize=8)
    for x_value, y_value in zip(nodes, q4):
        axes[0].annotate(f"{y_value:.3f}", (x_value, y_value), xytext=(0, -14), textcoords="offset points", ha="center", fontsize=8)
    axes[0].set_xscale("log", base=2)
    axes[0].set_xticks(nodes, [str(int(value)) for value in nodes])
    axes[0].set(title="网格加密下的烘干时间", xlabel="径向节点数", ylabel="连续事件时间 / h")
    axes[0].legend()
    clean_axis(axes[0])

    spacing = 1.0 / (nodes - 1.0)
    axes[1].loglog(spacing[:-1], e3[:-1], color=COLORS["blue"], marker="o", markersize=6, label="问题 3 · 相对 641 节点")
    axes[1].loglog(spacing[:-1], e4[:-1], color=COLORS["orange"], marker="s", markersize=6, label="问题 4 · 相对 641 节点")
    for x_value, value in zip(spacing[:-1], e3[:-1]):
        axes[1].annotate(f"{value:.1f}s", (x_value, value), xytext=(5, 5), textcoords="offset points", fontsize=8)
    for x_value, value in zip(spacing[:-1], e4[:-1]):
        axes[1].annotate(f"{value:.1f}s", (x_value, value), xytext=(5, -12), textcoords="offset points", fontsize=8)
    axes[1].set(
        title="空间离散误差衰减",
        xlabel="无量纲网格宽度 Δx",
        ylabel="相对 641 节点的时间差 / s",
    )
    axes[1].text(
        0.98, 0.05,
        f"Richardson：$p_3$={spatial['problem3_richardson']['observed_order']:.2f}，"
        f"$p_4$={spatial['problem4_richardson']['observed_order']:.2f}\n"
        f"估计 641 节点剩余误差：Q3 {spatial['problem3_richardson']['estimated_641_error_s']:.1f}s，"
        f"Q4 {spatial['problem4_richardson']['estimated_641_error_s']:.1f}s",
        transform=axes[1].transAxes, fontsize=8.5, va="bottom", ha="right",
        bbox={"boxstyle": "round,pad=0.45", "facecolor": "white", "edgecolor": COLORS["grid"], "alpha": 0.92},
    )
    axes[1].legend()
    clean_axis(axes[1])
    for letter, axis in zip("AB", axes):
        panel_label(axis, letter)
    figure.suptitle("网格收敛与空间离散误差", fontsize=16, fontweight="bold", color=COLORS["navy"])
    save_figure(figure, "04_grid_convergence_and_error")


def plot_diagnostics(data: dict[str, dict], error: dict) -> dict:
    environment = load_environment_xlsx(ROOT / "附件1.xlsx")
    mask = environment.time_s >= environment.stable_start_s
    stable_time = environment.time_s[mask] / 3600.0
    figure, axes = plt.subplots(2, 2, figsize=(13.2, 8.3), constrained_layout=True)

    axes[0, 0].plot(stable_time, environment.temperature_c[mask], color=COLORS["blue"], marker="o", markersize=3, label="实测")
    axes[0, 0].axhline(environment.stable_temperature_c, color=COLORS["red"], linewidth=2, label="稳定段均值")
    axes[0, 0].fill_between(stable_time, environment.stable_temperature_c - np.std(environment.temperature_c[mask], ddof=1), environment.stable_temperature_c + np.std(environment.temperature_c[mask], ddof=1), color=COLORS["blue"], alpha=0.13, label="±1σ")
    axes[0, 0].set(title="稳定阶段环境温度", xlabel="时间 / h", ylabel="温度 / °C")
    axes[0, 0].legend(ncol=3)
    clean_axis(axes[0, 0])

    axes[0, 1].plot(stable_time, environment.moisture[mask], color=COLORS["teal"], marker="o", markersize=3, label="实测")
    axes[0, 1].axhline(environment.stable_moisture, color=COLORS["orange"], linewidth=2, label="稳定段均值")
    axes[0, 1].fill_between(stable_time, environment.stable_moisture - np.std(environment.moisture[mask], ddof=1), environment.stable_moisture + np.std(environment.moisture[mask], ddof=1), color=COLORS["teal"], alpha=0.15, label="±1σ")
    axes[0, 1].set(title="稳定阶段环境含水率", xlabel="时间 / h", ylabel="含水率 / (kg/kg)")
    axes[0, 1].legend(ncol=3)
    clean_axis(axes[0, 1])

    q3 = np.asarray(data["result3"]["moisture"], dtype=float)
    q4 = np.asarray(data["result4"]["moisture"], dtype=float)
    t3 = np.asarray(data["result3"]["time"], dtype=float) / 3600.0
    t4 = np.asarray(data["result4"]["time"], dtype=float) / 3600.0
    axes[1, 0].plot(t3, q3[:, 0] - q3[:, -1], color=COLORS["blue"], label="问题 3：中心−表面")
    axes[1, 0].plot(t4, q4[:, 0] - q4[:, -1], color=COLORS["orange"], label="问题 4：中心−表面")
    axes[1, 0].axhline(0.0, color=COLORS["ink"], linewidth=1.0)
    axes[1, 0].set(title="径向含水率梯度的合理性", xlabel="时间 / h", ylabel="中心与表面含水率差")
    axes[1, 0].legend()
    clean_axis(axes[1, 0])

    sensitivity = error["environment_sensitivity"]
    labels = ["问题 3", "问题 4"]
    base = np.array([sensitivity["problem3"]["baseline_time_h"], sensitivity["problem4"]["baseline_time_h"]])
    lower = np.array([sensitivity["problem3"]["favorable_time_h"], sensitivity["problem4"]["favorable_time_h"]])
    upper = np.array([sensitivity["problem3"]["adverse_time_h"], sensitivity["problem4"]["adverse_time_h"]])
    x = np.arange(2)
    lower_minutes = (lower - base) * 60.0
    upper_minutes = (upper - base) * 60.0
    axes[1, 1].errorbar(
        x, np.zeros_like(base), yerr=np.vstack((-lower_minutes, upper_minutes)),
        fmt="o", color=COLORS["purple"], ecolor=COLORS["gold"],
        elinewidth=5, capsize=7, markersize=8, label="稳定环境 ±1σ 包络",
    )
    for index, value in enumerate(base):
        axes[1, 1].text(index, upper_minutes[index] + 0.8, f"基准 {value:.3f} h", ha="center", fontsize=8.5)
    axes[1, 1].axhline(0.0, color=COLORS["ink"], linewidth=1.0)
    axes[1, 1].set_xticks(x, labels)
    axes[1, 1].set(title="稳定平台波动对结束时间的敏感性", ylabel="相对基准时间变化 / min")
    axes[1, 1].legend()
    clean_axis(axes[1, 1])

    for letter, axis in zip("ABCD", axes.flat):
        panel_label(axis, letter)
    figure.suptitle("输入稳定性与物理解诊断", fontsize=16, fontweight="bold", color=COLORS["navy"])
    save_figure(figure, "05_reasonableness_and_sensitivity")

    finite_fields = []
    for name, field in [
        ("result1", "temperature"), ("result1", "moisture"),
        ("result2", "temperature"), ("result2", "moisture"),
        ("result3", "moisture"),
    ]:
        finite_fields.append(np.all(np.isfinite(np.asarray(data[name][field], dtype=float))))
    # 问题 4 的 NaN 是收缩后固定坐标在药材外部的预期空值，但 Inf 不允许。
    q4_all = np.asarray(data["result4"]["moisture"], dtype=float)
    finite_fields.append(not np.any(np.isinf(q4_all)))
    diagnostics = {
        "initial_temperature_max_abs_error_c": float(max(
            np.max(np.abs(np.asarray(data["result1"]["temperature"])[0] - 28.0)),
            np.max(np.abs(np.asarray(data["result2"]["temperature"])[0] - 28.0)),
        )),
        "initial_moisture_max_abs_error": float(max(
            np.max(np.abs(np.asarray(data["result1"]["moisture"])[0] - 2.55)),
            np.max(np.abs(np.asarray(data["result2"]["moisture"])[0] - 2.55)),
        )),
        "all_required_field_values_finite": bool(all(finite_fields)),
        "problem4_expected_outside_blank_count": int(np.count_nonzero(np.isnan(q4_all[:, :-1]))),
        "minimum_moisture": float(min(
            np.nanmin(np.asarray(data[name]["moisture"], dtype=float))
            for name in ("result1", "result2", "result3", "result4")
        )),
        "problem1_temperature_range_c": [
            float(np.min(np.asarray(data["result1"]["temperature"], dtype=float))),
            float(np.max(np.asarray(data["result1"]["temperature"], dtype=float))),
        ],
        "problem2_temperature_range_c": [
            float(np.min(np.asarray(data["result2"]["temperature"], dtype=float))),
            float(np.max(np.asarray(data["result2"]["temperature"], dtype=float))),
        ],
        "problem3_center_ge_surface": bool(np.all(q3[:, 0] >= q3[:, -1] - 1e-4)),
        "problem4_center_ge_surface": bool(np.all(q4[:, 0] >= q4[:, -1] - 1e-4)),
        "problem3_center_nonincreasing_after_4h": bool(np.all(np.diff(q3[t3 >= 4.0, 0]) <= 1e-4)),
        "problem4_center_nonincreasing_after_4h": bool(np.all(np.diff(q4[t4 >= 4.0, 0]) <= 1e-4)),
        "problem3_last_regular_center_moisture": float(q3[-1, 0]),
        "problem4_last_regular_center_moisture": float(q4[-1, 0]),
    }
    with (RESULTS / "reasonableness_metrics.json").open(
        "w", encoding="utf-8", newline="\n"
    ) as handle:
        json.dump(diagnostics, handle, ensure_ascii=False, indent=2)
    return diagnostics


def main() -> None:
    configure_style()
    data = load_payloads()
    summary = json.loads((INTERMEDIATE / "summary.json").read_text(encoding="utf-8"))
    error = json.loads((RESULTS / "error_analysis.json").read_text(encoding="utf-8"))
    plot_short_term_profiles(data)
    plot_long_term_comparison(data, summary)
    plot_spatiotemporal_fields(data)
    plot_grid_error(error)
    diagnostics = plot_diagnostics(data, error)
    print(f"已生成 5 组 PNG/SVG 图至 {FIGURES}")
    print(f"合理性检查：{json.dumps(diagnostics, ensure_ascii=False)}")


if __name__ == "__main__":
    main()
